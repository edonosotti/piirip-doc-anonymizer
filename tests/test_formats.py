import csv
from pathlib import Path
from zipfile import ZipFile

import pytest

from piirip.cli import main
from piirip.engine import Redactor
from piirip.playbooks import load_playbook

TEXT = "Contact John Smith at john.smith@example.com or +1 212-555-0198. IP 192.168.1.10."
SECRETS = ["john.smith@example.com", "John Smith", "192.168.1.10"]


def assert_package_redacted(path: Path, secret: str) -> None:
    with ZipFile(path) as package:
        assert all(secret.encode() not in package.read(name) for name in package.namelist())


@pytest.fixture(scope="session")
def redactor():
    return Redactor(load_playbook("default"))


def run(path: Path) -> Path:
    assert main([str(path)]) == 0
    return path.with_name(f"{path.stem}.redacted{path.suffix}")


def test_playbooks(redactor):
    assert "PERSON" in redactor.entities
    all_ = Redactor(load_playbook("all"))
    assert set(redactor.entities) < set(all_.entities)
    assert "IT_FISCAL_CODE" in all_.entities


def test_custom_playbook(tmp_path, redactor):
    pb = tmp_path / "pb.yaml"
    pb.write_text("entities: [EMAIL_ADDRESS]\n")
    r = Redactor(load_playbook(str(pb)))
    out = r.redact_text(TEXT)
    assert "<EMAIL_ADDRESS>" in out and "John Smith" in out


def test_txt_md_html(tmp_path):
    for name, content in [("a.txt", TEXT), ("a.md", f"# Hi\n{TEXT}"), ("a.html", f"<p>Hi <b>{TEXT}</b></p><a href='mailto:john.smith@example.com'>x</a>")]:
        f = tmp_path / name
        f.write_text(content)
        out = run(f).read_text()
        assert not any(s in out for s in SECRETS), out


def test_csv(tmp_path):
    f = tmp_path / "a.csv"
    with f.open("w", encoding="utf-8", newline="") as source:
        writer = csv.writer(source)
        writer.writerow(["id", "details", "notes"])
        writer.writerow(["1", TEXT, "Quoted, value\nwith another line"])
        writer.writerow(["2", "john.smith@example.com", "42"])

    with run(f).open(encoding="utf-8", newline="") as output:
        rows = list(csv.reader(output))

    assert len(rows) == 3
    assert rows[0] == ["id", "details", "notes"]
    assert rows[1][0] == "1"
    assert rows[1][2] == "Quoted, value\nwith another line"
    assert rows[2] == ["2", "<EMAIL_ADDRESS>", "42"]
    assert not any(secret in "\n".join(cell for row in rows for cell in row) for secret in SECRETS)


def test_docx_split_runs(tmp_path):
    import docx
    from docx.oxml import OxmlElement

    d = docx.Document()
    p = d.add_paragraph("Contact ")
    p.add_run("John ").bold = True
    p.add_run("Smith at john.smith@")
    p.add_run("example.com.")
    deleted = OxmlElement("w:del")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "john.smith@example.com"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    p._p.append(deleted)
    d.add_comment(
        p.runs, text="john.smith@example.com", author="john.smith@example.com"
    )
    d.add_table(1, 1).cell(0, 0).text = TEXT
    d.core_properties.author = "john.smith@example.com"
    f = tmp_path / "a.docx"
    d.save(f)
    output = run(f)
    out = docx.Document(str(output))
    text = "\n".join(p.text for p in out.paragraphs) + out.tables[0].cell(0, 0).text
    assert not any(s in text for s in SECRETS), text
    assert "<PERSON>" in text
    assert_package_redacted(output, "john.smith@example.com")
    with ZipFile(output) as package:
        assert b"&lt;EMAIL_ADDRESS&gt;" in package.read("word/comments.xml")


def test_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = TEXT
    wb.active["A2"] = 42
    wb.active["A1"].comment = openpyxl.comments.Comment(
        "john.smith@example.com", "john.smith@example.com"
    )
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "john.smith@example.com"
    wb.properties.title = "john.smith@example.com"
    f = tmp_path / "a.xlsx"
    wb.save(f)
    output = run(f)
    ws = openpyxl.load_workbook(output).active
    assert not any(s in ws["A1"].value for s in SECRETS)
    assert ws["A2"].value == 42
    sanitized = openpyxl.load_workbook(output)
    assert sanitized.active["A1"].comment.text == "<EMAIL_ADDRESS>"
    assert sanitized.active["A1"].comment.author in (None, "", "None")
    assert sanitized["Hidden"]["A1"].value == "<EMAIL_ADDRESS>"
    sanitized.close()
    assert_package_redacted(output, "john.smith@example.com")


def test_odt(tmp_path):
    from odf.opendocument import OpenDocumentText, load
    from odf.text import P, Span
    from odf.meta import InitialCreator
    from odf import teletype

    d = OpenDocumentText()
    p = P()
    p.addText("Contact ")
    s = Span()
    s.addText("John Smith")
    p.addElement(s)
    p.addText(" at john.smith@example.com. IP 192.168.1.10.")
    d.text.addElement(p)
    d.meta.addElement(InitialCreator(text="john.smith@example.com"))
    f = tmp_path / "a.odt"
    d.save(str(f))
    output = run(f)
    out = load(str(output))
    text = "".join(teletype.extractText(x) for x in out.text.childNodes)
    assert not any(s in text for s in SECRETS), text
    assert_package_redacted(output, "john.smith@example.com")


def test_pdf(tmp_path):
    import pymupdf as fitz

    d = fitz.open()
    d.new_page().insert_text((72, 72), TEXT)
    f = tmp_path / "a.pdf"
    d.save(f)
    out = fitz.open(str(run(f)))
    text = out[0].get_text()
    assert not any(s in text for s in SECRETS), text


def test_pdf_placeholder(tmp_path):
    import pymupdf as fitz

    d = fitz.open()
    d.new_page().insert_text((72, 72), TEXT)
    f = tmp_path / "a.pdf"
    d.save(f)
    assert main(["--pdf-style", "placeholder", str(f)]) == 0
    text = fitz.open(str(tmp_path / "a.redacted.pdf"))[0].get_text()
    assert not any(s in text for s in SECRETS), text
    assert "<" in text and "PERSON" in text


def test_pdf_hidden_content(tmp_path):
    import pymupdf as fitz

    d = fitz.open()
    page = d.new_page()
    page.insert_text((72, 72), "Visible text")
    annot = page.add_text_annot((72, 100), "john.smith@example.com")
    annot.set_info(title="john.smith@example.com")
    widget = fitz.Widget()
    widget.field_name = "john.smith@example.com"
    widget.field_label = "John Smith"
    widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    widget.field_value = "john.smith@example.com"
    widget.rect = fitz.Rect(72, 140, 220, 170)
    page.add_widget(widget)
    page.insert_link(
        {"kind": fitz.LINK_URI, "from": fitz.Rect(72, 200, 180, 220),
         "uri": "mailto:john.smith@example.com"}
    )
    d.set_metadata({"author": "john.smith@example.com"})
    d.set_xml_metadata(
        '<x:xmpmeta xmlns:x="adobe:ns:meta/">john.smith@example.com</x:xmpmeta>'
    )
    d.set_toc([[1, "john.smith@example.com", 1]])
    f = tmp_path / "a.pdf"
    d.save(f)
    d.close()

    output = run(f)
    out = fitz.open(str(output))
    assert out.metadata["author"] == ""
    assert out.get_xml_metadata() == ""
    assert out.get_toc()[0][1] == "<EMAIL_ADDRESS>"
    assert not any(
        "john.smith@example.com" in annot.info.values()
        for page in out for annot in (page.annots() or ())
    )
    assert not any(
        "john.smith@example.com" in (widget.field_name, widget.field_label, widget.field_value)
        for page in out for widget in (page.widgets() or ())
    )
    assert not any(
        "john.smith@example.com" in link.get("uri", "")
        for page in out for link in page.get_links()
    )
