"""Format handlers: each reads `src`, redacts PII and writes `dst`."""
from __future__ import annotations

from pathlib import Path
from typing import Callable

from .engine import Redactor


def _text(src: Path, dst: Path, r: Redactor) -> None:
    dst.write_text(r.redact_text(src.read_text(encoding="utf-8")), encoding="utf-8")


def _html(src: Path, dst: Path, r: Redactor) -> None:
    from bs4 import BeautifulSoup, Comment, NavigableString

    soup = BeautifulSoup(src.read_bytes(), "html.parser")
    skip = {"script", "style"}
    for node in list(soup.find_all(string=True)):
        if isinstance(node, Comment) or node.parent.name in skip:
            continue
        new = r.redact_text(str(node))
        if new != node:
            node.replace_with(NavigableString(new))
    for tag in soup.find_all(True):
        for attr in ("alt", "title", "href", "value", "content", "placeholder"):
            val = tag.get(attr)
            if isinstance(val, str) and val.strip():
                tag[attr] = r.redact_text(val)
    for c in soup.find_all(string=lambda s: isinstance(s, Comment)):
        c.extract()
    dst.write_text(str(soup), encoding="utf-8")


def _docx(src: Path, dst: Path, r: Redactor) -> None:
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(src))
    roots = [doc.element.body]
    for section in doc.sections:
        for part in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            roots.append(part._element)
    # Comments/footnotes live in separate parts.
    for rel in doc.part.rels.values():
        if rel.reltype.endswith(("/comments", "/footnotes", "/endnotes")) and hasattr(rel.target_part, "element"):
            roots.append(rel.target_part.element)
    seen = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for p in root.iter(qn("w:p")):
            ts = [t for t in p.iter(qn("w:t")) if _owner_paragraph(t, qn) is p]
            if ts:
                r.redact_segments(ts, lambda t: t.text, _set_wt)
    cp = doc.core_properties
    for attr in ("author", "last_modified_by", "comments", "title", "subject", "keywords"):
        setattr(cp, attr, "")
    doc.save(str(dst))


def _owner_paragraph(el, qn):
    """Nearest ancestor w:p (text boxes can nest paragraphs)."""
    parent = el.getparent()
    while parent is not None and parent.tag != qn("w:p"):
        parent = parent.getparent()
    return parent


def _set_wt(t, value: str) -> None:
    t.text = value
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _xlsx(src: Path, dst: Path, r: Redactor) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(src))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "s" and isinstance(cell.value, str):
                    cell.value = r.redact_text(cell.value)
                if cell.comment is not None:
                    cell.comment.text = r.redact_text(cell.comment.text)
                if cell.hyperlink is not None and cell.hyperlink.target:
                    cell.hyperlink.target = r.redact_text(cell.hyperlink.target)
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    wb.save(str(dst))


def _odt(src: Path, dst: Path, r: Redactor) -> None:
    from odf import teletype  # noqa: F401  (ensures odf text helpers are importable)
    from odf.opendocument import load
    from odf.namespaces import TEXTNS
    from odf.element import Text

    doc = load(str(src))

    def is_para(n) -> bool:
        return getattr(n, "qname", None) in ((TEXTNS, "p"), (TEXTNS, "h"))

    def leaves(node, out):
        for ch in node.childNodes:
            if ch.nodeType == ch.TEXT_NODE:
                out.append(ch)
            elif ch.nodeType == ch.ELEMENT_NODE and not is_para(ch):
                leaves(ch, out)

    def walk(node):
        for ch in list(node.childNodes):
            if ch.nodeType != ch.ELEMENT_NODE:
                continue
            if is_para(ch):
                out: list = []
                leaves(ch, out)
                if out:
                    r.redact_segments(out, lambda t: t.data, _set_odf)
                for sub in ch.childNodes:  # nested paragraphs (frames, notes)
                    if sub.nodeType == sub.ELEMENT_NODE:
                        walk(sub)
            else:
                walk(ch)

    def _set_odf(t, value):
        t.data = value

    walk(doc.text)
    for part in (doc.styles, doc.masterstyles):  # headers/footers
        walk(part)
    doc.save(str(dst))


def _placeholder_annot(page, rect, label: str, fitz) -> None:
    """White-out `rect` and write `label` in it, shrinking the font to fit."""
    size = max(min(rect.height * 0.8, 11.0), 4.0)
    if label:
        width = fitz.get_text_length(label, fontname="helv", fontsize=size)
        if width > rect.width:
            size = max(size * rect.width / width, 3.0)
    page.add_redact_annot(
        rect, text=label or None, fontname="helv", fontsize=size,
        fill=(1, 1, 1), text_color=(0, 0, 0), align=fitz.TEXT_ALIGN_LEFT,
    )


def _pdf(src: Path, dst: Path, r: Redactor) -> None:
    import pymupdf as fitz

    doc = fitz.open(str(src))
    for page in doc:
        text = page.get_text()
        spans = r.find(text)
        r.count += len(spans)
        done = set()
        for s in spans:
            frag = text[s.start : s.end].strip()
            if not frag or frag in done:
                continue
            done.add(frag)
            label = r.replacement(s)
            # Search each line separately so wrapped matches are still covered.
            first = True
            for line in filter(None, (l.strip() for l in frag.splitlines())):
                for rect in page.search_for(line):
                    if r.pdf_style == "box":
                        page.add_redact_annot(rect, fill=(0, 0, 0))
                    else:
                        _placeholder_annot(page, rect, label if first else "", fitz)
                first = False
        page.apply_redactions()
    doc.set_metadata({})
    doc.save(str(dst), garbage=4, deflate=True, clean=True)


HANDLERS: dict[str, Callable[[Path, Path, Redactor], None]] = {
    ".txt": _text,
    ".md": _text,
    ".markdown": _text,
    ".html": _html,
    ".htm": _html,
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".odt": _odt,
    ".pdf": _pdf,
}
